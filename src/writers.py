from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("match_score", "Match Score"),
    ("title", "Job Title"),
    ("company", "Company"),
    ("location", "Location"),
    ("matched_role", "Matched Target Role"),
    ("resume_id", "Resume To Use"),
    ("resume_label", "Resume Label"),
    ("posted_at", "Posted At (UTC)"),
    ("age_hours", "Age (hours)"),
    ("source", "Source"),
    ("matched_skills", "Matched Skills"),
    ("why_match", "Why Match"),
    ("url", "Apply URL"),
    ("apply_status", "Apply Status"),
    ("notes", "Your Notes"),
]


def write_excel(jobs: list[dict[str, Any]], path: Path, sheet_name: str = "Daily Matches") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    body_font = Font(name="Arial", size=10)
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, (_, header) in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, col, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row_i, job in enumerate(jobs, start=2):
        for col, (key, _) in enumerate(COLUMNS, start=1):
            val = job.get(key, "")
            if val is None:
                val = ""
            cell = ws.cell(row_i, col, val)
            cell.font = body_font
            cell.alignment = wrap

    widths = [12, 34, 22, 18, 28, 12, 22, 20, 12, 12, 28, 40, 40, 12, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(jobs) + 1)}"

    meta = wb.create_sheet("Run Info", 0)
    meta["A1"] = "Field"
    meta["B1"] = "Value"
    meta["A1"].font = header_font
    meta["B1"].font = header_font
    meta["A1"].fill = header_fill
    meta["B1"].fill = header_fill
    info = [
        ("Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        ("Candidate", "Athul S Babu"),
        ("Focus locations", "Dubai, Abu Dhabi, UAE"),
        ("Jobs listed", str(len(jobs))),
        ("How to use", "Set Apply Status to 'apply' for roles you want help drafting. Use Resume To Use column when submitting."),
        ("Assumption", "Sources are public job boards + optional RapidAPI JSearch. LinkedIn auto-apply is not used."),
    ]
    for i, (k, v) in enumerate(info, start=2):
        meta.cell(i, 1, k).font = body_font
        meta.cell(i, 2, v).font = body_font
    meta.column_dimensions["A"].width = 18
    meta.column_dimensions["B"].width = 90

    # move Daily Matches after Run Info — already created as active then Run Info at 0
    # Ensure Daily Matches exists
    wb.save(path)
    return path


def write_csv(jobs: list[dict[str, Any]], path: Path) -> Path:
    import csv

    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[c[0] for c in COLUMNS], extrasaction="ignore")
        writer.writeheader()
        for job in jobs:
            row = {k: job.get(k, "") for k, _ in COLUMNS}
            writer.writerow(row)
    return path


def write_json(jobs: list[dict[str, Any]], path: Path, meta: dict[str, Any] | None = None) -> Path:
    import json

    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "count": len(jobs),
        "meta": meta or {},
        "jobs": jobs,
    }
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return path
